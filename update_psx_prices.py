#!/usr/bin/env python3
"""
PSX snapshot price updater.

Fetches market-watch data from dps.psx.com.pk and writes an Excel workbook
with Live Prices and a compact PriceLookup sheet for XLOOKUP use.

Usage:
  pip install requests beautifulsoup4 pandas openpyxl lxml
  python update_psx_prices.py
  python update_psx_prices.py -o my_prices.xlsx -s HBL,UBL,OGDC
  python update_psx_prices.py --csv
"""

from __future__ import annotations

import argparse
import logging
import sys
from datetime import datetime
from pathlib import Path
from typing import Any

import pandas as pd
import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    datefmt="%H:%M:%S",
)
log = logging.getLogger("psx")

MARKET_WATCH_URL = "https://dps.psx.com.pk/market-watch"
DEFAULT_HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/122.0.0.0 Safari/537.36"
    ),
    "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
    "Referer": "https://dps.psx.com.pk/",
}
REQUEST_TIMEOUT = 30


class PSXFetchError(Exception):
    """Raised when market data cannot be retrieved or parsed."""


def _to_float(value: str) -> float | None:
    text = (value or "").replace(",", "").strip()
    if not text or text in {"-", "—", "N/A"}:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def fetch_market_watch(session: requests.Session) -> pd.DataFrame:
    """Download and parse the PSX market-watch table."""
    try:
        response = session.get(MARKET_WATCH_URL, timeout=REQUEST_TIMEOUT)
        response.raise_for_status()
    except requests.Timeout as exc:
        raise PSXFetchError("Request timed out while contacting PSX") from exc
    except requests.ConnectionError as exc:
        raise PSXFetchError("Connection failed — check network access") from exc
    except requests.HTTPError as exc:
        status = getattr(exc.response, "status_code", "?")
        raise PSXFetchError(f"HTTP error from PSX (status {status})") from exc
    except requests.RequestException as exc:
        raise PSXFetchError(f"Network error: {exc}") from exc

    soup = BeautifulSoup(response.text, "lxml")
    table = soup.find("table")
    if table is None:
        raise PSXFetchError(
            "Market-watch table not found — site layout may have changed"
        )

    rows: list[dict[str, Any]] = []
    for tr in table.find_all("tr")[1:]:
        cols = [td.get_text(strip=True) for td in tr.find_all("td")]
        if len(cols) < 10:
            continue
        try:
            rows.append(
                {
                    "Symbol": cols[0],
                    "Sector": cols[1],
                    "ListedIn": cols[2],
                    "LDCP": _to_float(cols[3]),
                    "Open": _to_float(cols[4]),
                    "High": _to_float(cols[5]),
                    "Low": _to_float(cols[6]),
                    "Current": _to_float(cols[7]),
                    "Change": _to_float(cols[8]),
                    "ChangePct": cols[9],
                    "Volume": cols[10].replace(",", "") if len(cols) > 10 else "",
                }
            )
        except (IndexError, TypeError):
            continue

    if not rows:
        raise PSXFetchError("No stock rows could be parsed from market-watch")

    return pd.DataFrame(rows)


def filter_symbols(df: pd.DataFrame, symbols: str) -> pd.DataFrame:
    if not symbols.strip():
        return df
    wanted = {s.strip().upper() for s in symbols.split(",") if s.strip()}
    filtered = df[df["Symbol"].str.upper().isin(wanted)].copy()
    if filtered.empty:
        raise PSXFetchError(f"None of the requested symbols were found: {wanted}")
    return filtered


def write_excel(df: pd.DataFrame, output: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Live Prices"

    title_font = Font(name="Arial", bold=True, size=14, color="FFFFFF")
    header_font = Font(name="Arial", bold=True, size=10, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E79")
    thin = Border(
        left=Side(style="thin", color="B0B0B0"),
        right=Side(style="thin", color="B0B0B0"),
        top=Side(style="thin", color="B0B0B0"),
        bottom=Side(style="thin", color="B0B0B0"),
    )
    center = Alignment(horizontal="center")

    ws.merge_cells("A1:K1")
    ws["A1"] = (
        f"PSX SNAPSHOT PRICES — Updated: "
        f"{datetime.now().strftime('%Y-%m-%d %H:%M:%S')} PKT"
    )
    ws["A1"].font = title_font
    ws["A1"].fill = header_fill
    ws["A1"].alignment = center

    ws.merge_cells("A2:K2")
    ws["A2"] = (
        "Source: dps.psx.com.pk/market-watch | "
        "Snapshot data (not tick-level real-time) | "
        "Verify on PSX / broker before trading"
    )
    ws["A2"].font = Font(name="Arial", italic=True, size=9, color="666666")

    headers = [
        "Symbol",
        "Sector",
        "LDCP",
        "Open",
        "High",
        "Low",
        "Current",
        "Change",
        "Change %",
        "Volume",
        "Listed In",
    ]
    for col, name in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col, value=name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = thin

    for i, row in enumerate(df.itertuples(index=False), start=5):
        values = [
            row.Symbol,
            row.Sector,
            row.LDCP,
            row.Open,
            row.High,
            row.Low,
            row.Current,
            row.Change,
            row.ChangePct,
            row.Volume,
            row.ListedIn,
        ]
        for col, value in enumerate(values, 1):
            cell = ws.cell(row=i, column=col, value=value)
            cell.border = thin
            cell.alignment = center
            if col in {3, 4, 5, 6, 7, 8}:
                cell.number_format = "#,##0.00"

    widths = {
        "A": 12,
        "B": 10,
        "C": 10,
        "D": 10,
        "E": 10,
        "F": 10,
        "G": 10,
        "H": 10,
        "I": 12,
        "J": 12,
        "K": 36,
    }
    for letter, width in widths.items():
        ws.column_dimensions[letter].width = width

    ws.freeze_panes = "A5"
    last_row = 4 + len(df)
    ws.auto_filter.ref = f"A4:K{last_row}"

    lookup = wb.create_sheet("PriceLookup")
    lookup["A1"] = "Symbol"
    lookup["B1"] = "Current"
    for col in (1, 2):
        cell = lookup.cell(row=1, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = thin

    for i, row in enumerate(df.itertuples(index=False), start=2):
        lookup.cell(row=i, column=1, value=row.Symbol).border = thin
        price_cell = lookup.cell(row=i, column=2, value=row.Current)
        price_cell.border = thin
        price_cell.number_format = "#,##0.00"

    lookup.column_dimensions["A"].width = 12
    lookup.column_dimensions["B"].width = 12

    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Fetch PSX market-watch snapshot and write Excel"
    )
    parser.add_argument(
        "-o",
        "--output",
        default="PSX_Live_Prices.xlsx",
        help="Output Excel path (default: PSX_Live_Prices.xlsx)",
    )
    parser.add_argument(
        "-s",
        "--symbols",
        default="",
        help="Comma-separated symbols to keep, e.g. HBL,UBL,OGDC",
    )
    parser.add_argument(
        "--csv",
        action="store_true",
        help="Also write a CSV next to the Excel file",
    )
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    output = Path(args.output)

    try:
        with requests.Session() as session:
            session.headers.update(DEFAULT_HEADERS)
            log.info("Fetching market-watch from PSX...")
            df = fetch_market_watch(session)

        log.info("Parsed %d symbols", len(df))
        df = filter_symbols(df, args.symbols)
        if args.symbols.strip():
            log.info("Filtered to %d symbols", len(df))

        write_excel(df, output)
        log.info("Saved Excel: %s", output.resolve())

        if args.csv:
            csv_path = output.with_suffix(".csv")
            df.to_csv(csv_path, index=False)
            log.info("Saved CSV: %s", csv_path.resolve())

        sample = df.head(8)[["Symbol", "Current", "ChangePct"]]
        log.info("Sample:\n%s", sample.to_string(index=False))
        return 0

    except PSXFetchError as exc:
        log.error("%s", exc)
        return 1
    except OSError as exc:
        log.error("File error while writing output: %s", exc)
        return 1
    except Exception as exc:
        log.exception("Unexpected error: %s", exc)
        return 1


if __name__ == "__main__":
    sys.exit(main())
